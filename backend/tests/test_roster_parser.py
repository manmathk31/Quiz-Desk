"""Unit tests for student roster CSV and XLSX parser (Phase 5)."""
import io
import openpyxl
import pytest

from app.services.roster_parser import (
    parse_and_validate_roster_file,
    generate_roster_csv_template,
    generate_roster_xlsx_template,
    MAX_ROSTER_ROWS,
)


def test_parse_valid_csv_standard_utf8():
    csv_content = (
        "prn,name,email,division\n"
        "202401110001,Aarav Sharma,Aarav.Sharma@College.EDU,A\n"
        "202401110002,Diya Patel,diya.patel@college.edu,B\n"
    ).encode("utf-8")

    result = parse_and_validate_roster_file(csv_content, filename="roster.csv")
    assert result.is_valid is True
    assert result.total_rows == 2
    assert result.valid_rows == 2
    assert len(result.errors) == 0

    assert result.rows[0].prn == "202401110001"
    assert result.rows[0].name == "Aarav Sharma"
    assert result.rows[0].email == "aarav.sharma@college.edu"  # Normalized lowercase
    assert result.rows[0].division == "A"

    assert result.rows[1].prn == "202401110002"
    assert result.rows[1].email == "diya.patel@college.edu"


def test_parse_valid_csv_with_utf8_bom():
    csv_text = (
        "prn,name,email,division\n"
        "202401110010,Sanjay Verma,sanjay@college.edu,C\n"
    )
    bom_content = b"\xef\xbb\xbf" + csv_text.encode("utf-8")

    result = parse_and_validate_roster_file(bom_content, filename="roster_bom.csv")
    assert result.is_valid is True
    assert result.total_rows == 1
    assert result.rows[0].prn == "202401110010"
    assert result.rows[0].email == "sanjay@college.edu"


def test_leading_zero_prn_preserved():
    csv_content = (
        "prn,name,email,division\n"
        "00123,Zero Lead Student,zero@college.edu,A\n"
        "000045,Another Lead,lead@college.edu,A\n"
    ).encode("utf-8")

    result = parse_and_validate_roster_file(csv_content, filename="leads.csv")
    assert result.is_valid is True
    assert result.rows[0].prn == "00123"
    assert result.rows[1].prn == "000045"


def test_missing_required_headers():
    csv_content = "prn,name,division\n001,John Doe,A\n".encode("utf-8")
    result = parse_and_validate_roster_file(csv_content, filename="bad_headers.csv")
    assert result.is_valid is False
    assert any("Missing required columns" in e.message for e in result.errors)


def test_empty_csv_and_header_only():
    empty_res = parse_and_validate_roster_file(b"", filename="empty.csv")
    assert empty_res.is_valid is False

    header_only = parse_and_validate_roster_file(
        b"prn,name,email,division\n", filename="header_only.csv"
    )
    assert header_only.is_valid is False
    assert any("no student data rows" in e.message for e in header_only.errors)


def test_duplicate_prn_in_same_file():
    csv_content = (
        "prn,name,email,division\n"
        "202401,Student One,one@college.edu,A\n"
        "202401,Student Duplicate,two@college.edu,A\n"
    ).encode("utf-8")

    result = parse_and_validate_roster_file(csv_content, filename="dup_prn.csv")
    assert result.is_valid is False
    assert any("Duplicate PRN '202401'" in e.message for e in result.errors)


def test_duplicate_email_in_same_file():
    csv_content = (
        "prn,name,email,division\n"
        "202401,Student One,common@college.edu,A\n"
        "202402,Student Two,COMMON@college.edu,B\n"
    ).encode("utf-8")

    result = parse_and_validate_roster_file(csv_content, filename="dup_email.csv")
    assert result.is_valid is False
    assert any("Duplicate email 'COMMON@college.edu'" in e.message for e in result.errors)


def test_invalid_email_format():
    csv_content = (
        "prn,name,email,division\n"
        "202401,Student One,not-an-email,A\n"
    ).encode("utf-8")

    result = parse_and_validate_roster_file(csv_content, filename="invalid_email.csv")
    assert result.is_valid is False
    assert any("Invalid email format" in e.message for e in result.errors)


def test_row_limit_enforcement():
    lines = ["prn,name,email,division"]
    for i in range(MAX_ROSTER_ROWS + 10):
        lines.append(f"PRN{i},Name {i},student{i}@college.edu,A")
    oversized = "\n".join(lines).encode("utf-8")

    result = parse_and_validate_roster_file(oversized, filename="huge.csv")
    assert result.is_valid is False
    assert any("exceeds the limit" in e.message for e in result.errors)


def test_parse_valid_xlsx_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["prn", "name", "email", "division"])
    ws.append(["007788", "Excel Student", "excel.student@college.edu", "Division-X"])
    ws.append(["202499", "Second Student", "second@college.edu", "Division-Y"])

    buf = io.BytesIO()
    wb.save(buf)
    content = buf.getvalue()

    result = parse_and_validate_roster_file(content, filename="roster.xlsx")
    assert result.is_valid is True
    assert result.total_rows == 2
    assert result.rows[0].prn == "007788"
    assert result.rows[0].name == "Excel Student"
    assert result.rows[0].email == "excel.student@college.edu"
    assert result.rows[0].division == "Division-X"


def test_template_generation():
    csv_tpl = generate_roster_csv_template()
    assert "prn,name,email,division" in csv_tpl
    assert "202401110001" in csv_tpl

    xlsx_tpl = generate_roster_xlsx_template()
    assert len(xlsx_tpl) > 0
    # Read back with openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_tpl))
    ws = wb.active
    assert ws.cell(row=1, column=1).value == "prn"
    assert ws.cell(row=1, column=2).value == "name"
    assert ws.cell(row=1, column=3).value == "email"
    assert ws.cell(row=1, column=4).value == "division"
