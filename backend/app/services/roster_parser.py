"""CSV and XLSX spreadsheet parsing and validation engine for Student Roster imports (DEC-007, Phase 5)."""
from dataclasses import dataclass
import csv
import io
import re
from typing import List, Optional, Set, Tuple

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# Bounded import limits
MAX_ROSTER_FILE_SIZE_BYTES = 2 * 1024 * 1024  # 2MB
MAX_ROSTER_ROWS = 500

REQUIRED_ROSTER_HEADERS = ["prn", "name", "email", "division"]

EMAIL_REGEX = re.compile(r"^[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+$")


@dataclass
class ParsedRosterRow:
    """Student candidate row extracted from roster spreadsheet."""
    row_number: int
    prn: str
    name: str
    email: str
    division: str


@dataclass
class RosterRowError:
    """Actionable validation error with row and column location."""
    row_number: int
    column: Optional[str]
    message: str


@dataclass
class RosterValidationResult:
    """Overall outcome of parsing and validating an uploaded roster file."""
    is_valid: bool
    total_rows: int
    valid_rows: int
    errors: List[RosterRowError]
    rows: List[ParsedRosterRow]


def is_xlsx_file(filename: Optional[str], content: bytes) -> bool:
    """Detect if file is an XLSX workbook via extension or PK zip signature."""
    if filename and filename.lower().endswith((".xlsx", ".xlsm")):
        return True
    return content.startswith(b"PK\x03\x04")


def parse_and_validate_roster_file(
    content: bytes,
    filename: Optional[str] = None,
) -> RosterValidationResult:
    """
    Parse and validate a student roster spreadsheet (.csv or .xlsx).
    Guarantees:
    - Bounded file size (<= 2MB)
    - Row count limit (<= 500 rows)
    - Required headers: prn, name, email, division
    - Preserves PRNs as strings (with leading zeros)
    - Normalizes emails to lowercase
    - Detects intra-file duplicate PRNs and emails
    - Returns actionable row-level errors
    """
    if len(content) > MAX_ROSTER_FILE_SIZE_BYTES:
        return RosterValidationResult(
            is_valid=False,
            total_rows=0,
            valid_rows=0,
            errors=[
                RosterRowError(
                    row_number=0,
                    column=None,
                    message=f"File exceeds maximum allowed size of {MAX_ROSTER_FILE_SIZE_BYTES // (1024 * 1024)}MB.",
                )
            ],
            rows=[],
        )

    if is_xlsx_file(filename, content):
        return _parse_xlsx_roster(content)
    else:
        return _parse_csv_roster(content)


def _parse_csv_roster(content: bytes) -> RosterValidationResult:
    """Parse CSV content with UTF-8 / UTF-8-BOM support."""
    try:
        decoded_text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            decoded_text = content.decode("latin-1")
        except Exception:
            return RosterValidationResult(
                is_valid=False,
                total_rows=0,
                valid_rows=0,
                errors=[
                    RosterRowError(
                        row_number=0,
                        column=None,
                        message="File encoding not recognized. Please upload a UTF-8 encoded CSV file.",
                    )
                ],
                rows=[],
            )

    reader = csv.reader(io.StringIO(decoded_text))
    raw_rows: List[List[str]] = []
    for r in reader:
        # Ignore completely empty rows
        if any(cell.strip() for cell in r):
            raw_rows.append(r)

    if not raw_rows:
        return RosterValidationResult(
            is_valid=False,
            total_rows=0,
            valid_rows=0,
            errors=[
                RosterRowError(
                    row_number=0,
                    column=None,
                    message="CSV file is empty.",
                )
            ],
            rows=[],
        )

    return _validate_roster_table(raw_rows)


def _parse_xlsx_roster(content: bytes) -> RosterValidationResult:
    """Parse Excel workbook content using openpyxl."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        return RosterValidationResult(
            is_valid=False,
            total_rows=0,
            valid_rows=0,
            errors=[
                RosterRowError(
                    row_number=0,
                    column=None,
                    message=f"Failed to read Excel workbook: {str(exc)}",
                )
            ],
            rows=[],
        )

    sheet = wb.active
    if not sheet:
        return RosterValidationResult(
            is_valid=False,
            total_rows=0,
            valid_rows=0,
            errors=[
                RosterRowError(
                    row_number=0,
                    column=None,
                    message="Workbook has no active sheets.",
                )
            ],
            rows=[],
        )

    raw_rows: List[List[str]] = []
    for row in sheet.iter_rows(values_only=True):
        if not row:
            continue
        row_values: List[str] = []
        has_content = False
        for cell in row:
            if cell is not None:
                val_str = str(cell).strip()
                if val_str:
                    has_content = True
                row_values.append(val_str)
            else:
                row_values.append("")
        if has_content:
            raw_rows.append(row_values)

    if not raw_rows:
        return RosterValidationResult(
            is_valid=False,
            total_rows=0,
            valid_rows=0,
            errors=[
                RosterRowError(
                    row_number=0,
                    column=None,
                    message="Excel worksheet is empty.",
                )
            ],
            rows=[],
        )

    return _validate_roster_table(raw_rows)


def _validate_roster_table(raw_rows: List[List[str]]) -> RosterValidationResult:
    """
    Validate headers, data rows, duplicates, and column formats.
    First row is treated as headers.
    """
    header_row = [h.strip().lower() for h in raw_rows[0]]
    header_indices = {}
    for idx, h in enumerate(header_row):
        if h in REQUIRED_ROSTER_HEADERS and h not in header_indices:
            header_indices[h] = idx

    missing_headers = [h for h in REQUIRED_ROSTER_HEADERS if h not in header_indices]
    if missing_headers:
        return RosterValidationResult(
            is_valid=False,
            total_rows=0,
            valid_rows=0,
            errors=[
                RosterRowError(
                    row_number=1,
                    column=None,
                    message=f"Missing required columns in header: {', '.join(missing_headers)}. Required columns are: {', '.join(REQUIRED_ROSTER_HEADERS)}.",
                )
            ],
            rows=[],
        )

    data_rows = raw_rows[1:]
    if not data_rows:
        return RosterValidationResult(
            is_valid=False,
            total_rows=0,
            valid_rows=0,
            errors=[
                RosterRowError(
                    row_number=1,
                    column=None,
                    message="File contains headers but no student data rows.",
                )
            ],
            rows=[],
        )

    if len(data_rows) > MAX_ROSTER_ROWS:
        return RosterValidationResult(
            is_valid=False,
            total_rows=len(data_rows),
            valid_rows=0,
            errors=[
                RosterRowError(
                    row_number=MAX_ROSTER_ROWS + 2,
                    column=None,
                    message=f"File contains {len(data_rows)} rows, which exceeds the limit of {MAX_ROSTER_ROWS} students per import.",
                )
            ],
            rows=[],
        )

    errors: List[RosterRowError] = []
    parsed_rows: List[ParsedRosterRow] = []

    seen_prns: dict[str, int] = {}
    seen_emails: dict[str, int] = {}

    for idx, row in enumerate(data_rows):
        row_num = idx + 2  # 1-based, header is row 1
        row_has_error = False

        def get_val(col_name: str) -> str:
            col_idx = header_indices[col_name]
            if col_idx < len(row):
                return str(row[col_idx]).strip()
            return ""

        prn_val = get_val("prn")
        name_val = get_val("name")
        email_val = get_val("email")
        division_val = get_val("division")

        # 1. PRN Validation
        if not prn_val:
            errors.append(RosterRowError(row_number=row_num, column="prn", message="PRN is required."))
            row_has_error = True
        elif len(prn_val) > 64:
            errors.append(RosterRowError(row_number=row_num, column="prn", message="PRN exceeds maximum length of 64 characters."))
            row_has_error = True
        else:
            norm_prn = prn_val.upper()
            if norm_prn in seen_prns:
                first_row = seen_prns[norm_prn]
                errors.append(
                    RosterRowError(
                        row_number=row_num,
                        column="prn",
                        message=f"Duplicate PRN '{prn_val}' detected in file (first appeared in row {first_row}).",
                    )
                )
                row_has_error = True
            else:
                seen_prns[norm_prn] = row_num

        # 2. Name Validation
        if not name_val:
            errors.append(RosterRowError(row_number=row_num, column="name", message="Student name is required."))
            row_has_error = True
        elif len(name_val) > 255:
            errors.append(RosterRowError(row_number=row_num, column="name", message="Student name exceeds maximum length of 255 characters."))
            row_has_error = True

        # 3. Email Validation
        if not email_val:
            errors.append(RosterRowError(row_number=row_num, column="email", message="Email address is required."))
            row_has_error = True
        else:
            norm_email = email_val.lower()
            if not EMAIL_REGEX.match(norm_email):
                errors.append(RosterRowError(row_number=row_num, column="email", message=f"Invalid email format: '{email_val}'."))
                row_has_error = True
            elif len(norm_email) > 255:
                errors.append(RosterRowError(row_number=row_num, column="email", message="Email address exceeds maximum length of 255 characters."))
                row_has_error = True
            elif norm_email in seen_emails:
                first_row = seen_emails[norm_email]
                errors.append(
                    RosterRowError(
                        row_number=row_num,
                        column="email",
                        message=f"Duplicate email '{email_val}' detected in file (first appeared in row {first_row}).",
                    )
                )
                row_has_error = True
            else:
                seen_emails[norm_email] = row_num

        # 4. Division Validation
        if not division_val:
            errors.append(RosterRowError(row_number=row_num, column="division", message="Division is required."))
            row_has_error = True
        elif len(division_val) > 32:
            errors.append(RosterRowError(row_number=row_num, column="division", message="Division exceeds maximum length of 32 characters."))
            row_has_error = True

        if not row_has_error:
            parsed_rows.append(
                ParsedRosterRow(
                    row_number=row_num,
                    prn=prn_val,
                    name=name_val,
                    email=email_val.lower(),
                    division=division_val,
                )
            )

    return RosterValidationResult(
        is_valid=(len(errors) == 0),
        total_rows=len(data_rows),
        valid_rows=len(parsed_rows),
        errors=errors,
        rows=parsed_rows,
    )


def generate_roster_csv_template() -> str:
    """Generate CSV template with required headers and representative sample rows."""
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(REQUIRED_ROSTER_HEADERS)
    writer.writerow(["202401110001", "Aarav Sharma", "aarav.sharma@college.edu", "A"])
    writer.writerow(["202401110002", "Diya Patel", "diya.patel@college.edu", "A"])
    writer.writerow(["202401110003", "Rohan Kulkarni", "rohan.kulkarni@college.edu", "B"])
    writer.writerow(["00124", "Ananya Rao", "ananya.rao@college.edu", "B"])
    return output.getvalue()


def generate_roster_xlsx_template() -> bytes:
    """Generate genuine XLSX template with styled header row and text-formatted PRN column."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Student Roster"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")

    headers = ["prn", "name", "email", "division"]
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    sample_data = [
        ["202401110001", "Aarav Sharma", "aarav.sharma@college.edu", "A"],
        ["202401110002", "Diya Patel", "diya.patel@college.edu", "A"],
        ["202401110003", "Rohan Kulkarni", "rohan.kulkarni@college.edu", "B"],
        ["00124", "Ananya Rao", "ananya.rao@college.edu", "B"],
    ]

    for row_idx, data in enumerate(sample_data, start=2):
        ws.append(data)
        # Ensure PRN column is stored explicitly as Text to preserve leading zeros
        prn_cell = ws.cell(row=row_idx, column=1)
        prn_cell.number_format = "@"

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 14

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
